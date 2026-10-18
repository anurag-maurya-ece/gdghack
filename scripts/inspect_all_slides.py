import pptx

prs = pptx.Presentation(r"C:\Users\HP\Downloads\Blue and Red Dynamic News Report Presentation.pptx")
print(f"Canvas: {prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f} inches")

for i, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {i+1} ====================")
    for j, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            tf = shape.text_frame
            full_text = tf.text.replace("\n", " \\n ")
            font_info = ""
            if tf.paragraphs and tf.paragraphs[0].runs:
                r0 = tf.paragraphs[0].runs[0]
                font_info = f"Font: {r0.font.name}, Size: {r0.font.size}, Bold: {r0.font.bold}"
            print(f"  Shape [{j}] '{shape.name}' Pos: ({shape.left/914400:.2f}\", {shape.top/914400:.2f}\", {shape.width/914400:.2f}\", {shape.height/914400:.2f}\")")
            print(f"    Text: \"{full_text}\" | {font_info}")
