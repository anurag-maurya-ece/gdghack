import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

TEMPLATE_PATH = r"C:\Users\HP\Downloads\Blue and Red Dynamic News Report Presentation.pptx"
BACKUP_PATH = r"C:\Users\HP\Downloads\Blue and Red Dynamic News Report Presentation_BACKUP.pptx"

# Create backup first
if not os.path.exists(BACKUP_PATH):
    import shutil
    shutil.copyfile(TEMPLATE_PATH, BACKUP_PATH)
    print("Backup created at:", BACKUP_PATH)

prs = pptx.Presentation(TEMPLATE_PATH)
print("Loaded successfully. Number of slides:", len(prs.slides))
