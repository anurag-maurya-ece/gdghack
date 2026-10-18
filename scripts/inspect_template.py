import pptx

prs = pptx.Presentation(r"C:\Users\HP\Downloads\Blue and Red Dynamic News Report Presentation.pptx")
print(f"Total Slides: {len(prs.slides)}")
print(f"Slide Size: {prs.slide_width} x {prs.slide_height}")

for i, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {i+1} ====================")
    for j, shape in enumerate(slide.shapes):
        txt = ""
        if shape.has_text_frame:
            txt = shape.text_frame.text.replace("\n", " ").strip()
            if len(txt) > 100:
                txt = txt[:100] + "..."
        print(f"  [{j}] Type: {shape.shape_type} | Name: {shape.name} | Pos: ({shape.left}, {shape.top}, {shape.width}, {shape.height}) | Text: '{txt}'")
